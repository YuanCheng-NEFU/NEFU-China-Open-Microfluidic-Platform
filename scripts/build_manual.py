"""Build the editable English manual and its Markdown companion.

Requirements: python-docx, Pillow and reportlab. Run from any working directory.
"""
from pathlib import Path
import json
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Manual'
FONT = 'Segoe UI'
INK = '000000'
GRAY = '526066'
LINE = 'D9D9D9'
TEAL = '006B63'


def shade(cell, fill):
    item = OxmlElement('w:shd')
    item.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(item)


def text(p, value, size=None, bold=None, color=None):
    for index, segment in enumerate(value.split('`')):
        run = p.add_run(segment)
        run.font.name = 'Consolas' if index % 2 else FONT
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def paragraph(doc, value, style=None):
    p = doc.add_paragraph(style=style)
    return text(p, value)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), LINE)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, value in [('top', '90'), ('bottom', '90'), ('left', '110'), ('right', '110')]:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:w'), value)
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    props.append(margins)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, values in enumerate([headers] + rows):
        row = t.rows[0] if i == 0 else t.add_row()
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, '263238' if i == 0 else ('F3F6F6' if i % 2 else 'FFFFFF'))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            text(p, str(value), size=9.5, bold=(i == 0), color='FFFFFF' if i == 0 else INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = Pt(3)
    return t


def picture(doc, name, caption, width=6.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    source = ROOT / name if '/' in name else ROOT / 'Hardware/Figures' / name
    run.add_picture(str(source), width=Inches(width))
    for prop in run._r.xpath('.//wp:docPr'):
        prop.set('descr', caption)
    cap = paragraph(doc, caption, 'Caption')
    cap.paragraph_format.space_after = Pt(12)


def add_field(p, instruction):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    p._p.append(field)


def build_pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak, KeepTogether

    regular, bold = 'Helvetica', 'Helvetica-Bold'
    fonts = Path('C:/Windows/Fonts')
    if (fonts / 'segoeui.ttf').is_file() and (fonts / 'segoeuib.ttf').is_file():
        pdfmetrics.registerFont(TTFont('Segoe', str(fonts / 'segoeui.ttf')))
        pdfmetrics.registerFont(TTFont('SegoeBold', str(fonts / 'segoeuib.ttf')))
        regular, bold = 'Segoe', 'SegoeBold'
    styles = {
        'body': ParagraphStyle('body', fontName=regular, fontSize=10.5, leading=12.4, spaceAfter=7),
        'title': ParagraphStyle('title', fontName=bold, fontSize=34, leading=39, spaceAfter=15),
        'subtitle': ParagraphStyle('subtitle', fontName=regular, fontSize=16, leading=20, spaceAfter=23),
        'heading': ParagraphStyle('heading', fontName=bold, fontSize=24, leading=29, spaceAfter=14),
        'subheading': ParagraphStyle('subheading', fontName=bold, fontSize=12.5, leading=15, spaceBefore=8, spaceAfter=7),
        'part': ParagraphStyle('part', fontName=bold, fontSize=9, leading=12, spaceAfter=7),
        'caption': ParagraphStyle('caption', fontName=regular, fontSize=9, leading=11, textColor=colors.HexColor('#526066'), spaceAfter=12),
        'cell': ParagraphStyle('cell', fontName=regular, fontSize=9.5, leading=11.4),
        'thead': ParagraphStyle('thead', fontName=bold, fontSize=9.5, leading=11.4, textColor=colors.white),
        'code': ParagraphStyle('code', fontName='Courier', fontSize=9, leading=11, leftIndent=11, spaceBefore=3, spaceAfter=9),
        'step': ParagraphStyle('step', fontName=regular, fontSize=10.5, leading=12.4, leftIndent=18, firstLineIndent=-18, spaceAfter=7),
    }
    def p(value, style='body'):
        return Paragraph(escape(str(value)).replace('\n', '<br/>').replace('`', ''), styles[style])
    def figure(name, caption, width):
        source = ROOT / name if '/' in name else ROOT / 'Hardware/Figures' / name
        item = Image(str(source))
        item.drawHeight *= width * inch / item.drawWidth
        item.drawWidth = width * inch
        return KeepTogether([item, Spacer(1, 3), p(caption, 'caption')])
    def page_frame(canvas, doc):
        canvas.setTitle('NEFU China Open Microfluidic Platform Build and Operation Guide')
        canvas.setAuthor('NEFU-China iDEC Experimental Group')
        if doc.page > 1:
            canvas.setFont(regular, 8)
            canvas.drawString(.82 * inch, 10.65 * inch, 'NEFU-CHINA    OPEN MICROFLUIDIC PLATFORM')
            canvas.setFillColor(colors.HexColor('#526066'))
            canvas.drawString(.82 * inch, .35 * inch, 'Build and Operation Guide     |     Edition ' + data['edition'])
            canvas.drawRightString(7.68 * inch, .35 * inch, str(doc.page))
    story = [p('NEFU-CHINA', 'subheading'), Spacer(1, 27), p('Open Microfluidic\nPlatform', 'title'),
             p('Build and Operation Guide', 'subtitle'), p('Total Control V7.0 PERF'),
             figure('Manual/images/microfluidic_preview.png', 'Microfluidic droplet observation from the optical-bench reference video.', 6),
             Spacer(1, 12), p('EDITION ' + data['edition'] + '     /     SEPTEMBER 2026', 'part'),
             p('NEFU-China iDEC Experimental Group\nNortheast Forestry University')]
    for page in data['pages']:
        story += [PageBreak(), p(page['part'], 'part'), p(page['title'], 'heading')]
        for item in page['content']:
            kind = item['type']
            if kind in ('p', 'h', 'code'):
                story.append(p(item['text'], {'p': 'body', 'h': 'subheading', 'code': 'code'}[kind]))
            elif kind == 'figure':
                story.append(figure(item['file'], item['caption'], item.get('width', 6.7)))
            elif kind == 'steps':
                story.extend(p(str(i) + '. ' + value, 'step') for i, value in enumerate(item['items'], 1))
            elif kind == 'table':
                rows = [[p(value, 'thead') for value in item['headers']]]
                rows.extend([p(value, 'cell') for value in row] for row in item['rows'])
                widths = item.get('widths', [6.8 / len(item['headers'])] * len(item['headers']))
                t = Table(rows, colWidths=[value * inch for value in widths], repeatRows=1, hAlign='CENTER')
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#263238')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F3F6F6'), colors.white]),
                    ('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#D9D9D9')),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('TOPPADDING', (0, 0), (-1, -1), 4.5), ('BOTTOMPADDING', (0, 0), (-1, -1), 4.5),
                    ('LEFTPADDING', (0, 0), (-1, -1), 5.5), ('RIGHTPADDING', (0, 0), (-1, -1), 5.5),
                ]))
                story += [t, Spacer(1, 6)]
    doc = SimpleDocTemplate(str(OUT / 'Build_and_Operation_Manual.pdf'), pagesize=(8.5 * inch, 11 * inch),
                            topMargin=.7 * inch, bottomMargin=.65 * inch, leftMargin=.82 * inch, rightMargin=.82 * inch)
    doc.build(story, onFirstPage=page_frame, onLaterPages=page_frame)


def main():
    data = json.loads((OUT / 'manual_content.json').read_text(encoding='utf-8'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.7), Inches(.65)
    section.left_margin, section.right_margin = Inches(.82), Inches(.82)
    section.header_distance, section.footer_distance = Inches(.28), Inches(.28)
    section.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 34), ('Subtitle', 16), ('Heading 1', 24), ('Heading 2', 12.5), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.font.bold = name != 'Subtitle'
        style.font.italic = False
        style.font.underline = False
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        for spacing in style.element.xpath('./w:rPr/w:spacing'):
            spacing.set(qn('w:val'), '0')
        style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    caption = doc.styles['Caption']
    caption.font.name = FONT
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    header = section.header.paragraphs[0]
    text(header, 'NEFU-CHINA    OPEN MICROFLUIDIC PLATFORM', size=8, color=INK)
    header.paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.clear_all()
    doc.styles['Footer'].paragraph_format.tab_stops.clear_all()
    text(footer, 'Build and Operation Guide     |     Edition ' + data['edition'], size=8, color=GRAY)
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run('\t')
    add_field(footer, 'PAGE')
    doc.core_properties.title = 'NEFU China Open Microfluidic Platform Build and Operation Guide'
    doc.core_properties.subject = 'Assembly, commissioning and operation of a fluorescence guided droplet sorting platform'
    doc.core_properties.author = 'NEFU-China iDEC Experimental Group'
    doc.core_properties.last_modified_by = 'NEFU-China iDEC Experimental Group'
    doc.core_properties.version = data['edition']
    doc.core_properties.comments = ''

    p = paragraph(doc, 'NEFU-CHINA', None)
    for run in p.runs:
        run.font.size = Pt(12)
        run.bold = True
    p.paragraph_format.space_after = Pt(37)
    paragraph(doc, 'Open Microfluidic\nPlatform', 'Title')
    p = paragraph(doc, 'Build and Operation Guide', 'Subtitle')
    p.paragraph_format.space_after = Pt(24)
    paragraph(doc, 'Total Control V7.0 PERF')
    picture(doc, 'Manual/images/microfluidic_preview.png', 'Microfluidic droplet observation from the optical-bench reference video.', width=6.0)
    p = paragraph(doc, 'EDITION ' + data['edition'] + '     /     SEPTEMBER 2026')
    p.paragraph_format.space_before = Pt(12)
    for r in p.runs:
        r.font.size = Pt(9)
        r.bold = True
    paragraph(doc, 'NEFU-China iDEC Experimental Group\nNortheast Forestry University')
    md = ['# NEFU China Open Microfluidic Platform', '', '## Build and Operation Guide', '', 'Edition ' + data['edition'] + ' | September 2026', '']

    for page in data['pages']:
        doc.add_page_break()
        p = paragraph(doc, page['part'])
        p.paragraph_format.space_after = Pt(5)
        for r in p.runs:
            r.font.size = Pt(9)
            r.bold = True
        paragraph(doc, page['title'], 'Heading 1')
        md += ['## ' + page['title'], '']
        for item in page['content']:
            kind = item['type']
            if kind == 'p':
                paragraph(doc, item['text'])
                md += [item['text'], '']
            elif kind == 'h':
                paragraph(doc, item['text'], 'Heading 2')
                md += ['### ' + item['text'], '']
            elif kind == 'steps':
                for index, step in enumerate(item['items'], 1):
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(.25)
                    p.paragraph_format.first_line_indent = Inches(-.25)
                    text(p, str(index) + '. ', bold=True)
                    text(p, step)
                    md += [str(index) + '. ' + step]
                md += ['']
            elif kind == 'table':
                table(doc, item['headers'], item['rows'], item.get('widths', [6.8 / len(item['headers'])] * len(item['headers'])))
                md += ['| ' + ' | '.join(item['headers']) + ' |', '| ' + ' | '.join(['---'] * len(item['headers'])) + ' |']
                md += ['| ' + ' | '.join(map(str, row)) + ' |' for row in item['rows']]
                md += ['']
            elif kind == 'figure':
                picture(doc, item['file'], item['caption'], item.get('width', 6.7))
                link = '../' + item['file'] if '/' in item['file'] else '../Hardware/Figures/' + item['file']
                md += ['![' + item['caption'] + '](' + link + ')', '']
            elif kind == 'code':
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(.16)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(9)
                p.paragraph_format.line_spacing = 1.05
                r = p.add_run(item['text'])
                r.font.name = 'Consolas'
                r.font.size = Pt(9)
                md += ['```' + item.get('lang', 'text'), item['text'], '```', '']
            else:
                raise ValueError(kind)
    path = OUT / 'Build_and_Operation_Manual.docx'
    doc.save(path)
    (OUT / 'Build_and_Operation_Manual.md').write_text('\n'.join(md), encoding='utf-8')
    build_pdf(data)
    print(path)


if __name__ == '__main__':
    main()
